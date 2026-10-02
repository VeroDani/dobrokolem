# Builds the PM review queue from the three Verifier v4 batches.
# Chief-of-staff pre-review: priority + note per record. Run from repo root.
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

RUN = "runs/2026-10-02-verifier-v4"
out = []
for b in (1, 2, 3):
    out += list(csv.DictReader(open(f"{RUN}/output-batch{b}.csv", encoding="utf-8")))
base = {r["id"]: r for r in csv.DictReader(open("data/p7-records-2026-10-01.csv", encoding="utf-8"))}

# Chief-of-staff notes: (priority, note). Priority: A sensitive, B disagreement with manual check, C conflict to decide, D routine
COS = {
    "P7-015": ("A", "Průvodce obsahuje soukromé mobilní číslo jednotlivce (R-101). Před jakýmkoli dalším použitím odstranit."),
    "P7-044": ("A", "Agent narazil na zmínku o volebním sporu (2022) – politicky citlivé. Rozhodni, zda a jak skupinu uvádět."),
    "P7-012": ("A", "Soukromý bytový dům – před zveřejněním na mapě potřeba souhlas (R-101)."),
    "P7-013": ("A", "Jmenovaný soukromý včelař – před zveřejněním souhlas (R-101)."),
    "P7-024": ("A", "Web MČ uvádí u koordinátorů soukromé mobily. Do databáze nepřebírat; kontakt jen přes FB skupinu nebo MČ."),
    "P7-020": ("B", "Agent pravděpodobně chybuje: MigAct na podzim 2026 vede s Agorou 7 a MČ vzdělávací program komunitních spojek na P7 (cokdekdy.cz, 9/2026) a web má © 2026. Návrh: verified_changed."),
    "P7-004": ("B", "Neshoda s ručním ověřením. Web Elpidy blokuje agenty; aktivitu potvrzuje Elpida/Přístav 7. Rychlá ruční kontrola."),
    "P7-014": ("B", "Neshoda s ručním ověřením. rcletna.cz blokuje agenty; výsledky hledání ukazují aktuální web s rozpisem herny. Rychlá ruční kontrola."),
    "P7-056": ("B", "Ruční ověření bylo mírnější (poslední ročník 30. 5. 2025 je už mimo 12 měsíců). Agent postupoval podle pravidla – zeptat se pořadatelů na ročník 2026."),
    "P7-063": ("B", "Organizace celopražsky aktivní, ale aktivitu přímo na P7 agent nedoložil. Agent je přísnější, ale konzistentní."),
    "P7-034": ("B", "Agent našel aktivitu, ruční ověření ne. Zkontroluj zdroj."),
    "P7-039": ("B", "Agent našel aktivitu, ruční ověření ne. Zkontroluj zdroj."),
    "P7-041": ("B", "Agent použil aktuální stránku služby MČ jako důkaz (sporné pravidlo – viz list Návrhy pravidel, N-1)."),
    "P7-042": ("B", "Agent našel aktuální čísla Hobuletu (2026). Pravděpodobně správně."),
    "P7-043": ("B", "Městský kanál – agent ověřil, ruční ověření ne. Pravděpodobně správně."),
    "P7-070": ("B", "Agent našel opravu knihobudek (únor–duben 2026) na praha7.cz. Pravděpodobně správně."),
    "P7-001": ("C", "Rozpor ve zdrojích (R-004): „nadace“ vs. nadační fond, kdo provozuje kurzy v Přístavu 7. Interně si ujasnit, jde o naši organizaci."),
    "P7-003": ("C", "Rozpor e-mailu: pristav@ vs. pristav7@elpida.cz – web Elpidy uvádí pristav@."),
    "P7-023": ("C", "Rozpor Šormová vs. Panko (R-004) – stejný jako u P7-029. Jedna odpověď radnice vyřeší oba."),
    "P7-057": ("C", "Rozpor pořadatele: NZM vs. Pro-Bio Liga (R-004)."),
    "P7-060": ("C", "Ověřeno jen podle nedatované stránky obchodu (sporné pravidlo – N-1). Dva weby provozovatele se liší u druhé prodejny."),
    "P7-022": ("C", "Odstavec převzatý z průvodce Prahou 3 (R-006). Agent navrhuje náhradu: bezplatný kurz češtiny MČ od 9/2026."),
    "P7-021": ("C", "Prázdný nadpis „(nutné doplnit)“, projekt nenalezen. Doplnit, nebo vyřadit."),
    "P7-030": ("C", "Useknutý e-mail v průvodci; poslední doložená aktivita 2019. Kandidát na výzvu e-mailem."),
    "P7-038": ("C", "Na webu spolku je stránka „Aktivity 2025/2026“, ale web blokuje agenty. 1 minuta ruční kontroly → verified."),
}
for i in ("P7-025", "P7-026", "P7-027", "P7-028"):
    COS[i] = ("A", COS["P7-024"][1])

def default_note(r):
    if r["status"] == "needs_check":
        return ("D", "K výzvě e-mailem / ruční kontrole (agent nenašel dost důkazů).")
    return ("D", "Rutinní – zkontroluj zjištění a zdroje.")

rows = []
for r in out:
    pri, note = COS.get(r["id"], default_note(r))
    rows.append((pri, r, note))
rows.sort(key=lambda x: (x[0], x[1]["id"]))

F = "Arial"
hfont = Font(name=F, bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", fgColor="2F5D50")
bfont = Font(name=F, size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D0D0")
yellow = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()
# Instructions
ws0 = wb.active
ws0.title = "Jak na to"
lines = [
    ("Fronta ke schválení · Ověřovatel v4 · 2. 10. 2026", True),
    ("Agent prověřil 51 záznamů. Na stupni autonomie 1 schvaluješ každý výstup.", False),
    ("", False),
    ("1. Zapiš čas začátku a konce kontroly (níže). Měříme tím north star agentů: živé záznamy na hodinu tvého času.", False),
    ("2. Jdi podle priority: A citlivé → B neshoda s ručním ověřením → C rozpory k rozhodnutí → D rutina.", False),
    ("3. Ve žlutých sloupcích vyber Rozhodnutí. Při „Opravit“ vyplň správný status a jednou větou proč – z toho se agent učí.", False),
    ("4. Záznamy D se stavem needs_check stačí hromadně schválit – půjdou do výzvy e-mailem.", False),
    ("", False),
    ("Začátek kontroly (hh:mm)", False),
    ("Konec kontroly (hh:mm)", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=13 if b else 10)
ws0["B9"].fill = yellow; ws0["B10"].fill = yellow
ws0["A12"] = "Souhrn"; ws0["A12"].font = Font(name=F, bold=True, size=11)
ws0.column_dimensions["A"].width = 95; ws0.column_dimensions["B"].width = 14

# Queue
ws = wb.create_sheet("Fronta")
heads = ["Priorita", "ID", "Název", "Status agenta", "Jistota", "Ruční ověření (1. 10.)", "Zjištění agenta",
         "Chybí", "Navržené změny", "Zdroje", "Poznámka pravé ruky", "Rozhodnutí PM", "Správný status", "Proč (1 věta)"]
widths = [8, 8, 30, 15, 9, 15, 50, 22, 36, 50, 50, 14, 16, 34]
ws.append(heads)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(row=1, column=i); c.font = hfont; c.fill = hfill
    c.alignment = Alignment(wrap_text=True, vertical="center")
for pri, r, note in rows:
    ws.append([pri, r["id"], base[r["id"]]["name"], r["status"], r["confidence"], base[r["id"]]["status"],
               r["finding_cs"], r["missing_fields"], r["proposed_changes"], r["evidence"], note, None, None, None])
last = ws.max_row
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row:
        c.font = bfont; c.alignment = wrap; c.border = Border(bottom=thin)
    for c in row[11:14]:
        c.fill = yellow
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:N{last}"
dv1 = DataValidation(type="list", formula1='"Schválit,Opravit,Zamítnout"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"verified,verified_changed,needs_check,closed,outside_area"', allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add(f"L2:L{last}"); dv2.add(f"M2:M{last}")
pri_col = {"A": "F4CCCC", "B": "FCE5CD", "C": "DDEBF7", "D": "EEEEEE"}
for p, col in pri_col.items():
    ws.conditional_formatting.add(f"A2:A{last}", FormulaRule(formula=[f'$A2="{p}"'], fill=PatternFill("solid", fgColor=col)))

# Summary formulas on sheet 1
summary = [
    ("Záznamů ve frontě", f"=COUNTA(Fronta!B2:B{last})"),
    ("– priorita A (citlivé)", f'=COUNTIF(Fronta!A2:A{last},"A")'),
    ("– priorita B (neshoda s ručním ověřením)", f'=COUNTIF(Fronta!A2:A{last},"B")'),
    ("– priorita C (rozpory)", f'=COUNTIF(Fronta!A2:A{last},"C")'),
    ("– priorita D (rutina)", f'=COUNTIF(Fronta!A2:A{last},"D")'),
    ("Rozhodnuto", f'=COUNTA(Fronta!L2:L{last})'),
    ("Opraveno", f'=COUNTIF(Fronta!L2:L{last},"Opravit")'),
    ("Míra oprav (cíl < 10 %)", f'=IF(B18=0,"",B19/B18)'),
    ("Doba kontroly (min)", '=IF(OR(B9="",B10=""),"",ROUND((B10-B9)*1440,0))'),
]
for i, (k, f) in enumerate(summary, 13):
    ws0.cell(row=i, column=1, value=k).font = bfont
    c = ws0.cell(row=i, column=2, value=f); c.font = Font(name=F, bold=True, size=10)
ws0["B20"].number_format = "0%"
ws0["B9"].number_format = "hh:mm"; ws0["B10"].number_format = "hh:mm"
ws0["A23"] = "Pozn.: časy zapiš ve tvaru 10:15."
ws0["A23"].font = Font(name=F, italic=True, size=9)

# Proposed rules
wr = wb.create_sheet("Návrhy pravidel")
wr.append(["#", "Problém (z dnešního běhu)", "Návrh", "Schválit? (ano/ne)"])
props = [
    ("N-1", "Dvě dávky narazily na rozpor v zadání: krok 6 říká, že nedatovaná stránka obchodu je důkaz se střední důvěrou, ale rozhodovací pravidlo pro „verified“ chce zdroj s vysokou důvěrou (P7-041, P7-060, P7-065, P7-067).",
     "Doplnit rozhodovací pravidlo: pro obchody, služby a městské instituce stačí aktuální vlastní stránka služby (adresa + otevírací doba) + jeden další zdroj se střední důvěrou."),
    ("N-2", "38 z 51 záznamů je needs_check. Hlavní příčina už není úsudek agenta, ale přístup: weby blokují automatické čtení (Elpida, rcletna.cz, sokol-bubenec.cz…) a 14 záznamů jsou čistě facebookové skupiny.",
     "Nepřidávat další pravidla pro agenta. needs_check posílat rovnou do výzvy e-mailem (Pisatel výzev) a pro blokované weby zavést 1minutovou ruční kontrolu ve frontě."),
    ("N-3", "Agenti počítají dotazy různě (dávka 2 započítala sdílené dotazy vícekrát; neúspěšná načtení se počítají).",
     "Sjednotit: lookups = počet volání WebSearch/WebFetch včetně neúspěšných; sdílený dotaz se počítá u prvního záznamu."),
    ("N-4", "Průvodce i web MČ zveřejňují soukromé mobily koordinátorů a jednotlivců (P7-015, P7-024–028).",
     "Nové pravidlo R-103: agent osobní kontakty nikdy nepřebírá ani ze zdánlivě „oficiálních“ zdrojů; nahlásí je PM jako riziko."),
]
for p in props:
    wr.append(list(p) + [None])
for i, w in enumerate([6, 70, 70, 16], 1):
    wr.column_dimensions[get_column_letter(i)].width = w
    c = wr.cell(row=1, column=i); c.font = hfont; c.fill = hfill
for row in wr.iter_rows(min_row=2):
    for c in row:
        c.font = bfont; c.alignment = wrap; c.border = Border(bottom=thin)
    row[3].fill = yellow
dv3 = DataValidation(type="list", formula1='"ano,ne"', allow_blank=True)
wr.add_data_validation(dv3); dv3.add("D2:D5")

wb.save(f"{RUN}/review-queue.xlsx")
print("rows", last - 1)
